"""Write docs/Agahi_Model_Guide.docx: a plain-English guide to every model in Agahi.
All results come from files in reports/ (made by training). Run after training:
    python scripts/make_model_doc.py
Needs python-docx (requirements-train.txt)."""
from __future__ import annotations

import csv
import json
import sys
from collections import Counter
from pathlib import Path

import _bootstrap  # noqa: F401
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from app.config import REPORTS_DIR, ROOT

OUT = ROOT / "docs" / "Agahi_Model_Guide.docx"
GREEN = RGBColor(0x1F, 0x3B, 0x29)
HORIZON_ORDER = ["d7", "d14", "d21", "d28", "m1", "m2", "m3"]
HORIZON_WORDS = {"d7": "7 days", "d14": "14 days", "d21": "21 days", "d28": "28 days",
                 "m1": "1 month", "m2": "2 months", "m3": "3 months"}


# ---------- loading results (reports/ only) ----------

def load_json(name: str) -> dict:
    p = REPORTS_DIR / name
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else {}


def load_csv(name: str) -> list[dict]:
    p = REPORTS_DIR / name
    if not p.exists():
        return []
    with open(p, encoding="utf-8") as f:
        return list(csv.DictReader(f))


def num(v) -> float | None:
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def pct(v, digits: int = 1) -> str:
    v = num(v)
    return "n/a" if v is None else f"{v * 100:.{digits}f}%"


# ---------- document helpers ----------

def shade(cell, hex_fill: str) -> None:
    tc = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc.append(shd)


def table(doc, header: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        run = c.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(c, "1F3B29")
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(v))
            run.font.size = Pt(9.5)
    if widths:
        from docx.shared import Cm
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()


def para(doc, text: str, bold_lead: str | None = None) -> None:
    p = doc.add_paragraph()
    if bold_lead:
        p.add_run(bold_lead).bold = True
    p.add_run(text)


def bullets(doc, items: list[str]) -> None:
    for it in items:
        doc.add_paragraph(it, style="List Bullet")


def heading(doc, text: str, level: int = 1) -> None:
    h = doc.add_heading(text, level=level)
    for r in h.runs:
        r.font.color.rgb = GREEN


# ---------- sections ----------

def overview(doc) -> None:
    heading(doc, "1. What Agahi does")
    para(doc, "Agahi is a text message (SMS) service for farmers and traders who sell vegetables at Kalimati "
              "wholesale market in Kathmandu. A farmer sends a short message from a basic phone, such as "
              "\"golbheda\" (tomato) or \"alu 2 hapta\" (potato in two weeks), and gets back today's market "
              "price, a price range for the coming days or months with an honest note on how sure it is, "
              "weather for their own district, and a simple tip on whether to sell now or keep the crop.")
    para(doc, "Every answer fits in at most two SMS. It works in English, Roman Nepali (Nepali typed with "
              "English letters) and Nepali script. No large language model runs when a farmer sends a "
              "message: everything is rules, word lists and small models trained on local data.")
    heading(doc, "How one message travels", 2)
    table(doc, ["Step", "What happens", "Where in the code"], [
        ["1", "Farmer sends an SMS to the Agahi number", "Twilio (or the web simulator)"],
        ["2", "Twilio forwards it to the app; the app checks Twilio's signature", "app/channels/sms_routes.py"],
        ["3", "The app finds the farmer's saved district, language and last question", "app/core/engine.py"],
        ["4", "The text is cleaned, spelling is fixed, crop, time and place are picked out", "app/core/nlu.py"],
        ["5", "A small classifier decides what the farmer wants (price, forecast, weather...)", "app/core/intent_lite.py"],
        ["6", "Prices, ready-made forecasts and weather are looked up", "app/core/data.py, app/weather/"],
        ["7", "A plain-language reply is filled in and trimmed to fit 2 SMS", "app/core/templates.py, replies.py"],
        ["8", "The reply goes back to Twilio, which sends it to the phone", "TwiML in the webhook response"],
    ], [1.2, 9.5, 5.5])
    heading(doc, "What runs when, and where", 2)
    table(doc, ["Part", "Runs when a message arrives?", "Runs only when training?"], [
        ["Cleaning text, spelling fixes, picking out crop, time and place", "Yes", "No"],
        ["Intent classifier (NumPy copy)", "Yes, about a millisecond", "Trained with scikit-learn"],
        ["Price forecasts", "No, read from the database", "Yes: backtests, model choice, final fit"],
        ["Weather days 1 to 7", "Yes, from Open-Meteo (cached 3 hours)", "No"],
        ["Weather weeks 2 to 4 and months 1 to 3", "Yes, from past-years tables", "Yes: tables and skill tests"],
        ["Reports, model card, this guide", "No", "Yes"],
    ], [7.5, 4.5, 4.5])


def prices(doc) -> None:
    heading(doc, "2. Price forecasting")
    para(doc, "Goal: for each crop, give a likely price range 7, 14, 21 and 28 days ahead, and a rough "
              "seasonal range 1, 2 and 3 months ahead, and say honestly how much to trust it.")
    heading(doc, "The data", 2)
    bullets(doc, [
        "Daily Kalimati wholesale prices per crop. For each crop we use the variety with the most days of data "
        "(for tomato, small local tomato).",
        "Days marked as carried over (the market repeated yesterday's price) are treated as missing. Gaps of up "
        "to 3 days are filled with the last price; longer gaps stay empty.",
        "Daily arrivals (tonnes reaching the market) and daily rain and temperature for the whole area.",
    ])
    heading(doc, "What is predicted", 2)
    para(doc, "We predict the change from today, not the price itself: the log of (future price divided by "
              "today's price). For 1 to 3 months we use the average of 7 days around the target day, because a "
              "single day that far ahead is mostly noise. The change is then turned back into rupees.")
    heading(doc, "Inputs (features)", 2)
    para(doc, "Every input for day t uses only information known on day t. Tests change the future data and check "
              "that the inputs for day t do not move.")
    table(doc, ["Input", "Plain meaning"], [
        ["Price changes over 1, 3, 7, 14 and 28 days", "Is the price rising or falling lately?"],
        ["Price against its 90-day average", "Is today unusually high or low?"],
        ["Price wobble over 14 days", "How jumpy has the price been?"],
        ["Arrivals against the last 30 days, and 7-day vs 28-day arrivals", "Is more or less produce coming in?"],
        ["Rain totals and rainy days over the last 3 to 28 days", "Did recent rain disturb picking or transport?"],
        ["Average top temperature over 7 and 28 days", "Heat stress and season"],
        ["Day of year and month", "Season"],
        ["Gap to the second variety (for example Indian vs local)", "Is one source much cheaper?"],
    ], [7.5, 8.5])
    para(doc, "Weather forecasts are not used as inputs, because there is no archive of past forecasts to test them on.")
    heading(doc, "The six candidate models", 2)
    table(doc, ["Model", "What it does", "Why it is there"], [
        ["Persistence", "Says the price will stay where it is today.", "The simplest honest guess. Any model must beat it."],
        ["Seasonal naive", "Uses the typical change at the same time of year in earlier years.", "Captures season without any fitting."],
        ["Mean reversion", "Expects the price to drift back toward its 90-day average.", "Prices that spike tend to come back."],
        ["Ridge regression", "A straight-line model on all inputs, kept small so it does not over-fit.", "Fast, stable, easy to explain."],
        ["Gradient boosting (HistGradientBoosting, median)", "Many small decision trees (depth 3) added together.", "Catches bends and thresholds a line misses."],
        ["Shrinkage blend", "Best simple model plus a share (0 to 100%) of the boosting model's correction.", "Takes the learned signal only as far as it helps."],
    ], [4.0, 6.5, 5.5])
    heading(doc, "How the models are tested: walk-forward with a gap", 2)
    para(doc, "We pretend to stand at many past dates. At each one, models are trained only on data before that "
              "date and asked about the next block of days (60 days in fast mode, 30 in full mode). The first "
              "training window is at least one year. Between the end of training and the start of the test block "
              "we leave a gap as long as the forecast horizon (the embargo), so that no training answer overlaps "
              "with the test period. The model with the smallest average error is chosen for each crop and horizon.")
    heading(doc, "Price ranges", 2)
    para(doc, "Ranges come from split-conformal calibration: we look at how wrong the chosen model was in earlier "
              "test blocks and widen the forecast by the 10th and 90th percentile of those errors. So the range is "
              "meant to contain the real price about 8 times in 10. Each test block is calibrated only on earlier blocks.")
    heading(doc, "Status rules and how to read them", 2)
    table(doc, ["Status", "Rule", "What the farmer reads (English)"], [
        ["Reliable", "Beats the best simple model by at least 5%, gets the direction right at least 58% of the time "
                     "(moves of 2% or more), range holds the real price 70% to 90% of the time, at least 60 test days.",
         "Good past record, but still a guess."],
        ["Indicative", "A learned model that is at least as good as the simple models, with a range that holds 70% to 90% of the time.",
         "Rough guess, not sure."],
        ["Pattern only", "Not better than the simple models.", "Same time last years: Rs.. to Rs.. This is not a forecast."],
        ["Unavailable", "Too little data, or the crop has had no new price for over a week.", "No forecast we can trust yet."],
    ], [2.6, 8.4, 5.0])
    para(doc, "Months 1 to 3 can never be reliable: there are just over three years of prices, so they are labelled "
              "as rough seasonal guesses. The sell or keep tip only uses 7 or 14 day forecasts that are reliable; "
              "otherwise it uses this week's price trend and tomorrow's rain at the farmer's district, and says "
              "there is no forecast to trust.")


def weather(doc) -> None:
    heading(doc, "3. Weather")
    table(doc, ["Time ahead", "Method", "What the reply says"], [
        ["Days 1 to 7", "Open-Meteo forecast for the farmer's district (free, no key). Cached 3 hours, 5 second timeout.",
         "Rain total, the wettest day, daytime temperature. If the internet is down: past-years averages, labelled."],
        ["Weeks 2 to 4", "Ridge models on recent rain and temperature compared with normal. Kept only if they beat "
                         "past-years averages by at least 3% in year-by-year tests with a 28-day gap.",
         "Usual weekly rain range, labelled \"from past years, not a forecast\" unless a model passed."],
        ["Months 1 to 3", "Past-years averages for the same 30 days, nudged by whether the last 30 days were wetter or drier than usual.",
         "Usual rain and whether the last 30 days were wetter or drier than usual."],
    ], [2.6, 7.4, 6.0])
    para(doc, "Past-years averages (climatology) are worked out for each district and each day of the year, "
              "smoothed over plus or minus 7 days, from 2015 onward.")


def nlp(doc, ev: dict) -> None:
    heading(doc, "4. Understanding messages (the chatbot)")
    bullets(doc, [
        "Word lists (aliases): every crop and district in English, Roman Nepali and Nepali script, with common "
        "misspellings (golbheda, golvheda, tamatar, गोलभेडा all mean tomato).",
        "Cleaning: lower case, Nepali digits to 0-9, merging spelling variants (ी and ि, ू and ु), cutting "
        "repeated letters (\"hooooo\" to \"hoo\"), expanding SMS shorthand (tmrw, kti, plz) and splitting Nepali "
        "word endings (आलुको to आलु को).",
        "Spelling fixes: an unknown word is matched against the word lists (RapidFuzz). It is only replaced when "
        "the best match is clearly better than any match with a different meaning, so common words are never "
        "turned into crops.",
        "Picking out details from one sentence: crop, how far ahead (14 din, 2 hapta, next month), district, "
        "quantity (500kg), variety (Indian, local) and time words (aaja, bholi, parsi, hijo, next week, yo hapta).",
        "Memory: short follow-ups reuse the last crop, time and question. \"and potato?\" asks the same about potato, "
        "\"ani?\" moves to the next step, \"aru?\" offers another crop, \"kati ho?\" asks the price of the last crop.",
        "Confidence: 0.75 or more, act. Between 0.45 and 0.75, reply with the top two guesses as a numbered choice. "
        "Below 0.45, show the main menu.",
        "Replies are fixed templates in plain words, one idea per line, with the numbered options on their own lines.",
    ])
    heading(doc, "The intent classifier", 2)
    comp = ev.get("comparison", {})
    para(doc, f"Features: {ev.get('features', 'TF-IDF word and character n-grams')}. Two classifiers were compared by "
              "cross-validation on the training split and the better one was kept.")
    table(doc, ["Classifier", "Cross-validation accuracy", "Spread", "Kept"],
          [[k, pct(v.get("cv_accuracy"), 2), pct(v.get("cv_std"), 2), "yes" if k == ev.get("chosen") else ""]
           for k, v in comp.items()] or [["not trained", "", "", ""]], [6.0, 4.0, 3.0, 2.0])
    para(doc, f"Training data: {ev.get('n_total', 'n/a')} generated examples, at least 200 per intent and language, "
              "made from realistic sentence patterns plus typos, dropped vowels, repeated letters, SMS shorthand, "
              "mixed English and Roman Nepali and filler words. The web app uses a NumPy copy of the trained model, "
              "so it does not need scikit-learn; a test checks both give the same answers.")
    heading(doc, "SMS budget and languages", 2)
    para(doc, "English and Roman Nepali use the GSM-7 alphabet: 160 characters in one SMS, 153 per part when longer. "
              "Nepali script uses UCS-2: 70 characters in one SMS, 67 per part. Line breaks count as characters. "
              "Every reply is held to 2 parts in code: the menu is laid out more tightly first, then the least "
              "important lines are dropped (menu, then the weather line, then the date note). The status line of a "
              "forecast is never dropped.")


def model_table(doc, fp: dict, ev: dict) -> None:
    heading(doc, "5. Every model at a glance")
    table(doc, ["Model", "Purpose", "Library", "Input", "Output", "File", "Size", "Time"], [
        ["Price models (91 crop x horizon cells)", "Price change range", "scikit-learn, NumPy",
         "Price, arrivals, weather features", "Change from today, P10/P50/P90", "models/price/*.joblib",
         f"{fp.get('price_bytes', 0) / 1024:.0f} KB total", f"{fp.get('median_inference_ms') or 0:.2f} ms median (training only)"],
        ["Intent classifier", "What the farmer wants", "scikit-learn to train, NumPy to run",
         "Cleaned message text", "Probability per intent", "models/intent/intent_lite.*",
         f"{fp.get('intent_bytes', 0) / 1024:.0f} KB", f"{num(ev.get('inference_ms')) or 0:.2f} ms"],
        ["Weekly weather models", "Weeks 2 to 4 rain and temperature", "scikit-learn to train, NumPy to run",
         "Recent anomalies, season", "Weekly value", "models/weather/*.joblib",
         f"{fp.get('weather_bytes', 0) / 1024:.0f} KB ({fp.get('weather_files', 0)} file{'s' if fp.get('weather_files', 0) != 1 else ''}, incl. past-years table)", "under 1 ms"],
        ["Past-years averages", "Normal weather per day of year", "pandas", "Weather history", "Mean and ranges",
         "models/weather/climatology.joblib", "included above", "under 1 ms"],
        ["Spelling and word lists", "Find crops, places, times", "RapidFuzz", "Message words", "Corrected words, slots",
         "app/core/aliases.py", "in code", "about 1 ms"],
    ], [2.6, 2.2, 2.2, 2.2, 2.2, 2.4, 1.6, 1.8])


def results(doc, rows: list[dict], wx: list[dict], ev: dict, fp: dict) -> None:
    heading(doc, "6. Results")
    synthetic = fp.get("is_sample")
    if synthetic:
        para(doc, "Every number in this section comes from SYNTHETIC sample data, not real prices.", "Warning: ")
    para(doc, "All numbers below are read from reports/ (backtest_price.csv, backtest_weather.csv, intent_eval.json, "
              "footprint.json), which training writes. Run scripts/export_report.py and scripts/make_model_doc.py "
              "after training to refresh them.")
    chosen = [r for r in rows if r.get("chosen") == "1"]
    counts = Counter(r["status"] for r in chosen)
    para(doc, f"{len(chosen)} crop and horizon cells: " + ", ".join(f"{k} {v}" for k, v in sorted(counts.items())) + ".",
         "Price forecasts. ")
    order = {h: i for i, h in enumerate(HORIZON_ORDER)}
    chosen.sort(key=lambda r: (r["crop"], order.get(r["horizon"], 9)))
    table(doc, ["Crop", "Horizon", "Model kept", "Status", "Better than simple models by", "Range held price",
                "Direction right"],
          [[r["crop"], HORIZON_WORDS.get(r["horizon"], r["horizon"]), r["model"], r["status"], pct(r["skill"]),
            pct(r["coverage"]), pct(r["dir_acc"])] for r in chosen] or [["no results yet", "", "", "", "", "", ""]],
          [2.4, 2.0, 2.6, 2.2, 2.6, 2.2, 2.0])
    kept = sum(1 for r in wx if r.get("kept") == "1")
    skills = [num(r["skill"]) for r in wx if num(r.get("skill")) is not None]
    para(doc, f"{kept} of {len(wx)} district, week and measure combinations beat past-years averages by 3% or more. "
              + (f"Skill ranged from {min(skills) * 100:.0f}% to {max(skills) * 100:.0f}%. " if skills else "")
              + "So weeks 2 to 4 are served as past-years averages.", "Weather weeks 2 to 4. ")
    para(doc, f"held-out accuracy {pct(ev.get('accuracy'))} overall ("
              + ", ".join(f"{k} {pct(v)}" for k, v in ev.get("by_lang", {}).items())
              + f"). On {ev.get('casual_n', 0)} hand-written casual and noisy messages, the whole understanding step "
                f"got {pct(ev.get('casual_pipeline_accuracy'))} right; the classifier alone got "
                f"{pct(ev.get('casual_classifier_accuracy'))}.", "Message understanding: ")
    misses = ev.get("casual_misses", [])
    if misses:
        para(doc, "Weak spots (casual messages it still gets wrong):")
        table(doc, ["Message", "Expected", "Got"], [[m["text"], m["expected"], m["got"]] for m in misses], [8.0, 4.0, 4.0])


def limits(doc) -> None:
    heading(doc, "7. Limits and trade-offs")
    bullets(doc, [
        "Prices are Kalimati wholesale prices, not what a farmer gets at the farm gate.",
        "Arrivals include Indian imports, so they show market supply, not local harvest.",
        "Just over three years of prices: months 1 to 3 are rough seasonal guesses, and many cells are honestly "
        "\"pattern only\".",
        "Weather beyond about two weeks is past-years averages; no weekly model beat them in our tests.",
        "District detail is limited to six weather points plus the area average.",
        "The intent classifier learned from generated sentences. Real farmers will write things it has not seen; "
        "the Quality queue collects those so they can be added.",
        "Nepali and Roman Nepali wording must be checked by a native speaker before a pilot.",
        "Twilio trial accounts add a prefix to every message, which uses part of the SMS budget.",
    ])


def glossary(doc) -> None:
    heading(doc, "8. Glossary")
    table(doc, ["Term", "Meaning"], [
        ["Baseline", "A very simple forecast (no change, same as last year, back to average) used as the bar to beat."],
        ["Skill", "How much smaller the error is than the best baseline. 10% skill means 10% less error."],
        ["Walk-forward test", "Testing a model on the past as if it were the future, one block at a time."],
        ["Embargo", "A gap between training data and test data so answers cannot leak."],
        ["P10, P50, P90", "Low, middle and high values: the real value should be below P10 1 time in 10 and above P90 1 time in 10."],
        ["Coverage", "How often the real value fell inside the range."],
        ["Climatology", "Past-years averages for this time of year."],
        ["Intent", "What the farmer wants: price, forecast, weather, sell or keep, arrivals, and so on."],
        ["Slot", "A detail picked out of the message: crop, time, district, quantity."],
        ["Segment", "One SMS part. Long messages are split into segments and each costs money."],
        ["GSM-7 / UCS-2", "The two SMS alphabets: Latin letters (160 per SMS) and everything else, such as Nepali (70 per SMS)."],
        ["TwiML", "The short XML reply Twilio reads to know what SMS to send back."],
    ], [4.0, 12.0])


def models_page(doc) -> None:
    heading(doc, "9. How to read the Models page")
    bullets(doc, [
        "The matrix shows every crop (rows) and horizon (columns). Green is reliable, amber indicative, grey pattern "
        "only, red unavailable. The number is skill against the best simple model.",
        "Click a cell to see real prices against the model's past predictions with the 80% range, error by horizon "
        "against the best simple model, how often the range held, the error spread, which inputs mattered and the "
        "live record of forecasts made since launch.",
        "The intent panel shows held-out accuracy by language, the two classifiers compared, and the casual-message "
        "results. The table under it lists the casual messages it still gets wrong.",
        "The footprint card shows the size of every model on disk and how long it takes to run on a normal CPU.",
    ])


def main() -> int:
    ev = load_json("intent_eval.json")
    fp = load_json("footprint.json")
    rows = load_csv("backtest_price.csv")
    wx = load_csv("backtest_weather.csv")
    if not rows:
        print("reports/backtest_price.csv not found. Train first (python scripts/train_all.py).")
        return 1
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    title = doc.add_heading("Agahi model guide", 0)
    for r in title.runs:
        r.font.color.rgb = GREEN
    p = doc.add_paragraph("How Agahi turns market and weather data into short, honest SMS answers. "
                          "Plain English first, technical detail where it is needed.")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    overview(doc)
    prices(doc)
    weather(doc)
    nlp(doc, ev)
    model_table(doc, fp, ev)
    results(doc, rows, wx, ev, fp)
    limits(doc)
    glossary(doc)
    models_page(doc)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Wrote {OUT}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
